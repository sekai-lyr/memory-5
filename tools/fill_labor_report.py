from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


BODY_FONT = "宋体"
BODY_SIZE = Pt(12)


def set_run_font(run, size=BODY_SIZE, bold=False):
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = size
    run.bold = bold


def format_para(paragraph, first_line=False, align=None, bold=False, size=BODY_SIZE):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_line:
        fmt.first_line_indent = Pt(24)
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def set_para_text(paragraph, text, first_line=False, bold=False, size=BODY_SIZE):
    paragraph.text = ""
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    format_para(paragraph, first_line=first_line, bold=bold, size=size)


def set_cell_text(cell, text):
    cell.text = ""
    for index, part in enumerate(text.split("\n")):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        set_para_text(paragraph, part)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_paragraph_after(paragraph, text, first_line=False, bold=False):
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p = new_element
    new_paragraph._element = new_element
    set_para_text(new_paragraph, text, first_line=first_line, bold=bold)
    return new_paragraph


def find_paragraph(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    return None


def fill_report():
    downloads = Path("C:/Users/administration/Downloads")
    source = next(downloads.glob("*大学生劳动教育*课程报告*docx"))
    output = Path("B:/Sekai/Sekai/sekai/memory-5/target/labor-education-report-filled.docx")
    doc = Document(str(source))

    table = doc.tables[0]
    set_cell_text(table.rows[0].cells[1], "（请填写）")
    set_cell_text(table.rows[0].cells[3], "（请填写）")
    set_cell_text(table.rows[0].cells[5], "（请填写）")
    set_cell_text(table.rows[1].cells[1], "个人实践")
    set_cell_text(table.rows[1].cells[3], "方案制定、工具准备、现场劳动、过程记录、报告撰写")
    set_cell_text(table.rows[2].cells[1], "如有团队成员请补充姓名")
    set_cell_text(table.rows[2].cells[3], "环境整理、资料归档、照片记录")
    set_cell_text(table.rows[3].cells[1], "如无团队可填“无”")
    set_cell_text(table.rows[3].cells[3], "协助检查、总结反思")
    set_cell_text(table.rows[4].cells[1], "校园环境维护与宿舍内务整理劳动实践")
    set_cell_text(table.rows[5].cells[1], "2026年5月（可按实际劳动日期修改）")
    set_cell_text(
        table.rows[6].cells[1],
        "本人在活动期间，严格要求自己，遵守国家法律法规、学校有关劳动实践活动的各项规章制度，按照学校的实践活动规定和实践活动计划完成好劳动实践活动。"
        "活动过程中，自觉维护学校形象，注意劳动安全和劳动保护，并对自己在活动期间的行为和安全负责；活动结束后，按时返校参加实践活动考核。"
        "我保证实践报告中的内容及信息尊重他人知识产权，所提交的业绩证据都是我自己及团队成员完成的，所有内容是真实有效的。\n学生签名：        2026 年 5 月 30 日",
    )
    set_cell_text(
        table.rows[7].cells[1],
        "本次劳动教育实践以校园环境维护和宿舍内务整理为主要内容，围绕“在劳动中理解责任、在实践中提升能力”的目标展开。"
        "活动前，我根据宿舍和公共区域的实际情况制订清扫计划，准备抹布、扫帚、垃圾袋等工具，并明确整理桌面、清理地面、分类废弃物、检查安全隐患等任务。"
        "实践过程中，我重点完成个人学习区域整理、宿舍公共地面清扫、可回收物分类、楼道卫生维护和劳动过程记录。活动结束后，宿舍环境更加整洁，物品摆放更有秩序，也增强了我的劳动意识和集体责任感。",
    )
    set_cell_text(
        table.rows[8].cells[1],
        "形成了较完整的劳动实践记录：一是完成宿舍桌面、地面、柜面和公共区域清扫；二是对废纸、塑料瓶、外包装等进行分类处理；"
        "三是拍摄整理前后对比照片并记录主要劳动步骤；四是总结出“先分类、后清洁、再复查”的行动方案。"
        "通过本次实践，宿舍卫生状况明显改善，个人物品归置更加清晰，公共空间使用效率得到提高。",
    )
    set_cell_text(
        table.rows[9].cells[1],
        "通过劳动实践，我认识到劳动不是简单重复的体力活动，而是培养责任感、秩序感和自我管理能力的重要方式。"
        "本次活动的特色在于把日常生活场景作为劳动课堂，将宿舍整理、垃圾分类和安全检查结合起来，使劳动成果能够直接服务学习生活。"
        "不足之处是前期分工和时间安排还不够细致，后续应坚持定期整理，形成长期习惯。",
    )
    for row_index in (7, 8, 9):
        shade_cell(table.rows[row_index].cells[0], "EAF4FF")

    sections = [
        (
            "一、活动概述：（简述实践活动的意义、内容、对象、准备、过程、效果等）",
            [
                "劳动教育是高校立德树人的重要环节，也是大学生认识劳动价值、提升实践能力、养成良好生活习惯的重要途径。通过参加校园环境维护与宿舍内务整理劳动，我进一步体会到整洁有序的学习生活环境并不是自然形成的，而是依靠每个人主动参与、持续维护和相互配合。此次实践活动以宿舍和校园公共区域为主要对象，围绕卫生清扫、物品整理、垃圾分类、安全检查和过程记录等内容展开，既贴近大学生日常生活，也具有较强的现实意义。",
                "活动开始前，我先对宿舍桌面、地面、床铺周边和公共区域进行了观察，发现主要问题包括物品摆放不够整齐、纸张和包装袋堆放较多、地面角落存在灰尘、部分公共物品缺少固定摆放位置等。针对这些问题，我制订了较为具体的劳动计划：第一步进行物品分类，把学习资料、生活用品、可回收物和废弃物分开；第二步清理桌面和柜面，擦拭容易积灰的位置；第三步清扫地面和公共区域，重点处理角落、门后和桌下；第四步检查插座、排插和杂物堆放情况，减少安全隐患；第五步拍摄整理前后照片，记录劳动过程和成果。",
                "实践过程中，我按照计划逐项完成任务。首先将桌面上的书籍、文具和电子设备重新归类摆放，把暂时不用的资料装入文件夹，减少桌面杂乱感。随后清理宿舍垃圾，将废纸、塑料瓶和普通生活垃圾分开处理，并把可回收物集中投放到指定位置。在公共区域清洁中，我重点清扫地面、擦拭桌椅和柜面，同时提醒自己注意劳动安全，避免湿手接触插座。劳动结束后，宿舍整体环境有了明显改善，桌面更加整洁，通道更加通畅，公共空间也更适合学习和休息。此次实践虽然内容普通，但让我感受到日常劳动对个人成长和集体生活的重要作用。",
            ],
        ),
        (
            "二、实践成果：（列举过程性成果：计划、过程记录、行动方案、相关图片等，）",
            [
                "本次实践形成了较为完整的过程性成果。首先是劳动计划成果：我根据宿舍实际情况确定了“分类整理、重点清洁、复查维护”的工作思路，并将任务划分为桌面整理、地面清扫、垃圾分类、公共区域维护和安全检查五个部分，使劳动过程更加有条理。其次是过程记录成果：在劳动过程中，我记录了清扫前存在的问题、使用的劳动工具、完成的主要步骤以及整理后的变化，便于后续总结经验。",
                "在行动方案方面，我总结出一套适合宿舍日常维护的小方法：一是每天保持桌面基本整洁，学习用品使用后及时归位；二是每周固定时间清理地面和角落，防止灰尘积累；三是将垃圾分类和宿舍值日结合起来，减少临时突击打扫；四是公共物品设置固定位置，避免重复寻找和随意堆放；五是定期检查插座、充电线和杂物摆放，保障宿舍安全。通过这些方法，劳动成果不只是一次性的整洁，而可以逐渐转化为长期习惯。",
                "相关图片说明如下：图片1为宿舍桌面整理前后对比，可体现学习资料归类和桌面空间改善情况；图片2为地面及公共区域清扫过程，可体现实际劳动环节；图片3为垃圾分类和可回收物整理情况，可体现环保意识与劳动实践结合；图片4为整理后的宿舍整体效果，可体现劳动成果。若提交纸质版或电子版，可在本报告相应位置插入真实照片，并在照片下方标注时间、地点和劳动内容。",
            ],
        ),
        (
            "三、活动感悟：（谈谈自己的活动收获、本次活动的特色及创新、存在的问题及反思等）",
            [
                "通过本次劳动教育实践，我对劳动的价值有了更具体的理解。过去我常把宿舍清扫、物品整理看作普通的日常事务，认为只要不影响生活即可。但真正投入实践后，我发现劳动不仅能改善环境，也能改变人的状态。一个整洁的空间会让学习更专注、生活更有秩序；一次认真完成的清扫，也能让人意识到公共环境需要每个人共同维护。劳动并不只在课堂之外，它本身就是一种学习方式，能够培养耐心、责任心和解决实际问题的能力。",
                "本次活动的特色在于把劳动教育与真实生活场景结合起来。相比单纯参加一次集中劳动，宿舍内务整理更贴近大学生的日常，也更容易检验劳动成果是否能够持续。活动中，我尝试用“问题清单+步骤记录+成果复查”的方式推进任务，先找出问题，再按步骤解决，最后检查效果。这种方式让劳动过程更清楚，也让我明白，做好一件小事同样需要计划意识和执行能力。",
                "当然，本次实践也存在不足。比如前期对时间的估计不够准确，部分角落清理耗时较长；照片记录还可以更规范，应该在整理前、中、后分别拍摄，形成更完整的证据链；后续维护机制也需要进一步坚持，不能只在提交报告前集中整理。今后我会把劳动实践融入日常生活，坚持定期整理宿舍、主动维护公共环境，并在劳动中不断提升自我管理能力。劳动教育让我认识到，大学生成长不仅体现在知识学习上，也体现在能否认真对待生活、承担责任、服务集体。",
            ],
        ),
    ]

    # Remove placeholder ellipsis paragraphs.
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip() in {"……", "..."}:
            delete_paragraph(paragraph)

    # Fill text directly under each existing attachment heading.
    for heading, body_paragraphs in sections:
        heading_para = find_paragraph(doc, heading)
        if heading_para is None:
            continue
        cursor = heading_para
        for text in body_paragraphs:
            cursor = insert_paragraph_after(cursor, text, first_line=True)

    photo_heading = doc.add_paragraph("图片附件粘贴处（请将真实劳动照片替换到下列位置）")
    format_para(photo_heading, bold=True)
    photo_table = doc.add_table(rows=2, cols=2)
    photo_table.style = "Table Grid"
    labels = [
        "图片1：整理前桌面与公共区域情况",
        "图片2：清扫地面、擦拭桌椅过程",
        "图片3：垃圾分类与可回收物整理",
        "图片4：整理完成后的宿舍环境",
    ]
    for index, cell in enumerate(photo_table._cells):
        set_cell_text(cell, labels[index] + "\n（在此处粘贴照片）")
        shade_cell(cell, "F7FBFF")

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            format_para(paragraph)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    format_para(paragraph)

    title = find_paragraph(doc, "金陵科技学院《大学生劳动教育》实践报告")
    if title is not None:
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            set_run_font(run, Pt(18), True)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    fill_report()
