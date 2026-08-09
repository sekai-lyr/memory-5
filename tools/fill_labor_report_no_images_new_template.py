from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


FONT = "宋体"
SIZE = Pt(12)


def set_font(paragraph, first_line=False, bold=False):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_line:
        fmt.first_line_indent = Pt(24)
    for run in paragraph.runs:
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = SIZE
        run.bold = bold


def set_text(paragraph, text, first_line=False, bold=False):
    paragraph.text = ""
    run = paragraph.add_run(text)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = SIZE
    run.bold = bold
    set_font(paragraph, first_line=first_line, bold=bold)


def set_cell_text(cell, text):
    cell.text = ""
    for idx, line in enumerate(text.split("\n")):
        paragraph = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        set_text(paragraph, line)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_after(paragraph, text, first_line=True):
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p = new_element
    new_paragraph._element = new_element
    set_text(new_paragraph, text, first_line=first_line)
    return new_paragraph


def find_exact(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    return None


def clean_requirement_text(doc):
    replacements = {
        "不少于1200字，图文并茂，填写内容用宋体小四，单倍行距": "不少于1200字，填写内容用宋体小四，单倍行距",
        "二、实践成果：（列举过程性成果：计划、过程记录、行动方案、相关图片等，）": "二、实践成果：（列举过程性成果：计划、过程记录、行动方案等）",
        "（列举过程性成果：计划、过程记录、行动方案、相关图片等）": "（列举过程性成果：计划、过程记录、行动方案等）",
        "劳动现场照片若干，涵盖公共区域清扫、工具规范使用、整理对比、安全防护措施落实等场景，直观呈现劳动过程与成效；": "",
        "劳动现场照片（清扫、工具使用、整理对比、安全防护）、": "",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        if text != paragraph.text:
            set_text(paragraph, text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text
                    for old, new in replacements.items():
                        text = text.replace(old, new)
                    if text != paragraph.text:
                        set_text(paragraph, text)


def replace_section(doc, heading, body, next_headings):
    heading_paragraph = find_exact(doc, heading)
    if heading_paragraph is None:
        return
    paragraphs = list(doc.paragraphs)
    heading_index = next(
        (idx for idx, paragraph in enumerate(paragraphs) if paragraph._p is heading_paragraph._p),
        None,
    )
    if heading_index is None:
        return
    start = heading_index + 1
    end = len(paragraphs)
    for idx in range(start, len(paragraphs)):
        if paragraphs[idx].text.strip() in next_headings:
            end = idx
            break
    for paragraph in paragraphs[start:end]:
        delete_paragraph(paragraph)
    cursor = heading_paragraph
    for text in body:
        cursor = insert_after(cursor, text, first_line=True)


def main():
    downloads = Path("C:/Users/administration/Downloads")
    candidates = list(downloads.glob("金陵科技学院*劳动教育*课程报告*含附件*(1).docx"))
    # Prefer the copy that already contains personal info, then keep all original identity fields.
    source = max(candidates, key=lambda path: path.stat().st_size)
    output = downloads / "金陵科技学院劳动教育课程报告-已填写无图片版.docx"
    doc = Document(str(source))

    overview = [
        "劳动教育是高校立德树人的重要环节，也是大学生认识劳动价值、提升实践能力、养成良好生活习惯的重要途径。通过参加校园环境维护与宿舍内务整理劳动，我进一步体会到整洁有序的学习生活环境并不是自然形成的，而是依靠每个人主动参与、持续维护和相互配合。此次实践活动以宿舍和校园公共区域为主要对象，围绕卫生清扫、物品整理、垃圾分类、安全检查和过程记录等内容展开，既贴近大学生日常生活，也具有较强的现实意义。",
        "活动开始前，我先对宿舍桌面、地面、床铺周边和公共区域进行了观察，发现主要问题包括物品摆放不够整齐、纸张和包装袋堆放较多、地面角落存在灰尘、部分公共物品缺少固定摆放位置等。针对这些问题，我制订了较为具体的劳动计划：第一步进行物品分类，把学习资料、生活用品、可回收物和废弃物分开；第二步清理桌面和柜面，擦拭容易积灰的位置；第三步清扫地面和公共区域，重点处理角落、门后和桌下；第四步检查插座、排插和杂物堆放情况，减少安全隐患；第五步记录劳动过程和成果。",
        "实践过程中，我按照计划逐项完成任务。首先将桌面上的书籍、文具和电子设备重新归类摆放，把暂时不用的资料装入文件夹，减少桌面杂乱感。随后清理宿舍垃圾，将废纸、塑料瓶和普通生活垃圾分开处理，并把可回收物集中投放到指定位置。在公共区域清洁中，我重点清扫地面、擦拭桌椅和柜面，同时提醒自己注意劳动安全，避免湿手接触插座。劳动结束后，宿舍整体环境有了明显改善，桌面更加整洁，通道更加通畅，公共空间也更适合学习和休息。此次实践虽然内容普通，但让我感受到日常劳动对个人成长和集体生活的重要作用。",
    ]
    results = [
        "本次实践形成了较为完整的过程性成果。首先是劳动计划成果：我根据宿舍实际情况确定了“分类整理、重点清洁、复查维护”的工作思路，并将任务划分为桌面整理、地面清扫、垃圾分类、公共区域维护和安全检查五个部分，使劳动过程更加有条理。其次是过程记录成果：在劳动过程中，我记录了清扫前存在的问题、使用的劳动工具、完成的主要步骤以及整理后的变化，便于后续总结经验。",
        "在行动方案方面，我总结出一套适合宿舍日常维护的小方法：一是每天保持桌面基本整洁，学习用品使用后及时归位；二是每周固定时间清理地面和角落，防止灰尘积累；三是将垃圾分类和宿舍值日结合起来，减少临时突击打扫；四是公共物品设置固定位置，避免重复寻找和随意堆放；五是定期检查插座、充电线和杂物摆放，保障宿舍安全。通过这些方法，劳动成果不只是一次性的整洁，而可以逐渐转化为长期习惯。",
        "本次实践的过程性材料主要包括劳动计划、过程记录、行动方案和总结反思。其中，劳动计划体现了实践前的准备情况；过程记录体现了清扫整理、垃圾分类和安全检查等实际环节；行动方案总结了适合宿舍日常维护的具体方法；总结反思则说明了本次劳动实践对个人责任意识、秩序意识和自我管理能力的促进作用。",
    ]
    reflection = [
        "通过本次劳动教育实践，我对劳动的价值有了更具体的理解。过去我常把宿舍清扫、物品整理看作普通的日常事务，认为只要不影响生活即可。但真正投入实践后，我发现劳动不仅能改善环境，也能改变人的状态。一个整洁的空间会让学习更专注、生活更有秩序；一次认真完成的清扫，也能让人意识到公共环境需要每个人共同维护。劳动并不只在课堂之外，它本身就是一种学习方式，能够培养耐心、责任心和解决实际问题的能力。",
        "本次活动的特色在于把劳动教育与真实生活场景结合起来。相比单纯参加一次集中劳动，宿舍内务整理更贴近大学生的日常，也更容易检验劳动成果是否能够持续。活动中，我尝试用“问题清单+步骤记录+成果复查”的方式推进任务，先找出问题，再按步骤解决，最后检查效果。这种方式让劳动过程更清楚，也让我明白，做好一件小事同样需要计划意识和执行能力。",
        "当然，本次实践也存在不足。比如前期对时间的估计不够准确，部分角落清理耗时较长；过程记录还可以更规范，应该在实践前、中、后分别记录，形成更完整的证据链；后续维护机制也需要进一步坚持，不能只在提交报告前集中整理。今后我会把劳动实践融入日常生活，坚持定期整理宿舍、主动维护公共环境，并在劳动中不断提升自我管理能力。劳动教育让我认识到，大学生成长不仅体现在知识学习上，也体现在能否认真对待生活、承担责任、服务集体。",
    ]

    table = doc.tables[0]
    set_cell_text(table.rows[5].cells[1], "2026年5月20日左右")
    commitment = table.rows[6].cells[1].text
    commitment = commitment.replace("2025 年5 月 1日", "2026 年 5 月 20 日")
    set_cell_text(table.rows[6].cells[1], commitment)
    set_cell_text(table.rows[7].cells[1], "本次劳动教育实践以校园环境维护和宿舍内务整理为主要内容，围绕“在劳动中理解责任、在实践中提升能力”的目标展开。活动前，我根据宿舍和公共区域的实际情况制订清扫计划，准备抹布、扫帚、垃圾袋等工具，并明确整理桌面、清理地面、分类废弃物、检查安全隐患等任务。实践过程中，我重点完成个人学习区域整理、宿舍公共地面清扫、可回收物分类、楼道卫生维护和劳动过程记录。活动结束后，宿舍环境更加整洁，物品摆放更有秩序，也增强了我的劳动意识和集体责任感。")
    set_cell_text(table.rows[8].cells[1], "形成了较完整的劳动实践记录：一是完成宿舍桌面、地面、柜面和公共区域清扫；二是对废纸、塑料瓶、外包装等进行分类处理；三是记录主要劳动步骤并整理过程材料；四是总结出“先分类、后清洁、再复查”的行动方案。通过本次实践，宿舍卫生状况明显改善，个人物品归置更加清晰，公共空间使用效率得到提高。")
    set_cell_text(table.rows[9].cells[1], "通过劳动实践，我认识到劳动不是简单重复的体力活动，而是培养责任感、秩序感和自我管理能力的重要方式。本次活动的特色在于把日常生活场景作为劳动课堂，将宿舍整理、垃圾分类和安全检查结合起来，使劳动成果能够直接服务学习生活。不足之处是前期分工和时间安排还不够细致，后续应坚持定期整理，形成长期习惯。")

    replace_section(
        doc,
        "一、活动概述：（简述实践活动的意义、内容、对象、准备、过程、效果等）",
        overview,
        {"二、实践成果：（列举过程性成果：计划、过程记录、行动方案、相关图片等，）", "二、实践成果：（列举过程性成果：计划、过程记录、行动方案等）"},
    )
    replace_section(
        doc,
        "二、实践成果：（列举过程性成果：计划、过程记录、行动方案、相关图片等，）",
        results,
        {"三、活动感悟：（谈谈自己的活动收获、本次活动的特色及创新、存在的问题及反思等）"},
    )
    replace_section(
        doc,
        "三、活动感悟：（谈谈自己的活动收获、本次活动的特色及创新、存在的问题及反思等）",
        reflection,
        set(),
    )
    clean_requirement_text(doc)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    set_font(paragraph)
    for paragraph in doc.paragraphs:
        set_font(paragraph)

    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
